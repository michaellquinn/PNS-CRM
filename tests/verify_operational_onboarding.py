"""Execute the operational workflow without mutating a real CRM database."""
import asyncio
import importlib.util
import io
import json
from pathlib import Path
import sys
import types
from datetime import datetime, timedelta

stub = types.ModuleType("asyncmy")
stub.cursors = types.SimpleNamespace(DictCursor=object)
sys.modules.setdefault("asyncmy", stub)
sys.modules.setdefault("asyncmy.cursors", stub.cursors)
spec = importlib.util.spec_from_file_location("operational_test_main", Path(__file__).parents[1] / "backend/main.py")
m = importlib.util.module_from_spec(spec)
spec.loader.exec_module(m)
NOW = datetime(2026, 9, 20, 19)
m.ob_now = lambda: NOW
sales = m.User(email="sales@example.test", name="Sales Owner", group="Commercial", level="staff")
manager = m.User(email="manager@example.test", name="Manager", group="Commercial", level="manager")
team = m.User(email="fleet@example.test", name="Fleet", group="4W", level="staff")
qc = m.User(email="qc@example.test", name="QC", group="QC", level="staff")
t = dict(id=1, ticket_ref="SOF-1", shipper="PT Test", service_type="FTL", opportunity_id="123",
         opportunity_name="PT Test - FTL", status="Proposal Accepted / Ready to Ship", sales_name=sales.name,
         sales_email=sales.email, potential_rev=999999, deleted_at=None)
p = {key: options[0] if options else "Operational detail" for key, _, _, _, options, _ in m.OB_FIELDS}
p.update(packing=["PCK", "PCK Kayu"], pickup_at="2026-09-21", planned_golive="2026-09-21",
         pickup_time="10-12", delivery_time="16-18", service="FTL", opportunity_id="123", global_id="G-1",
         shipper_id="8001", pickup_function="4W", delivery_function="2W", product_volume="20",
         pickup_tkbm="Yes", pickup_tkbm_count="2", delivery_tkbm="No", delivery_tkbm_count="0",
         implan="No", implan_count="0", rdo="No", rdo_treatment="Not required: shipper does not use RDO",
         pod_treatment="Return signed POD", surat_jalan_treatment="Return original Surat Jalan")
docs = [dict(kind="product_photo"), dict(kind="pickup_points")]

def rejected(call, code):
    try:
        call()
    except m.HTTPException as e:
        assert e.status_code == code, (e.status_code, e.detail)
    else:
        raise AssertionError("Expected request rejection")

async def denied(call, code):
    try:
        await call()
    except m.HTTPException as e:
        assert e.status_code == code, (e.status_code, e.detail)
    else:
        raise AssertionError("Expected request rejection")

assert m.ob_validate(p, t, docs) == datetime(2026, 9, 21, 10)
# Free-text detail is optional and uploads are no longer required (Michael, 2026-09-17).
assert m.ob_validate({**p, "pod_treatment": "", "handling": "", "pickup_driver": ""}, t, docs)
assert m.ob_validate(p, t, [])
# Hour ranges (Michael, 2026-09-17).
rejected(lambda: m.ob_validate({**p, "delivery_time": "18-16"}, t, docs), 400)
rejected(lambda: m.ob_validate({**p, "delivery_time": "16:00"}, t, docs), 400)
# A No switch sets its quantity to 0 automatically.
assert m.ob_zero_counts({"pickup_tkbm": "No", "pickup_tkbm_count": "", "implan": "Yes", "implan_count": "3"}) == \
    {"pickup_tkbm": "No", "pickup_tkbm_count": "0", "implan": "Yes", "implan_count": "3"}
# A field that follows the charter blocks submission when blank, and says where to fix it.
try:
    m.ob_validate({**p, "mps": ""}, t, docs)
except m.HTTPException as e:
    assert e.status_code == 400 and "Project Charter" in e.detail, e.detail
else:
    raise AssertionError("blank MPS must block submission")
rejected(lambda: m.ob_validate({**p, "pickup_address": ""}, t, docs), 400)
rejected(lambda: m.ob_validate({**p, "packing": ["No", "PCK"]}, t, docs), 400)
assert "sla" not in {k for k, *_ in m.OB_FIELDS}
assert set(m.OB_FOLLOW_CHARTER) == {"product_type", "delivery_mode", "mps", "rdo", "pickup_frequency"}
assert m.ob_source_value({"commodity": "", "product": "Snacks"}, ("commodity", "product")) == "Snacks"
assert all(f["section"] != "D · Pickup" for f in m.ob_schema())
rejected(lambda: m.ob_validate({**p, "pickup_at": "2026-09-20", "pickup_time": "23-24", "planned_golive": "2026-09-20"}, t, docs), 400)
# The first pickup is a date; its moment is that date at the start of the pickup range.
assert m.ob_pickup_moment(p) == datetime(2026, 9, 21, 10)
# Waiting time is a band, and an older charter's free text lands in one (2026-09-23).
assert m.ob_wait_bucket("2") == "1-2 hours" and m.ob_wait_bucket("0.5") == "< 1 hour"
assert m.ob_wait_bucket("3") == "2-3 hours" and m.ob_wait_bucket("4 hours") == "> 3 hours"
assert m.ob_wait_bucket("1-2 hours") == "1-2 hours"
assert m.ob_wait_bucket("None") == "" and m.ob_wait_bucket("") == ""
rejected(lambda: m.ob_validate({**p, "pickup_wait": "2 jam-ish"}, t, docs), 400)
# Admin may confirm for a team, and it is recorded as such (Michael, 2026-09-18).
_admin = m.User(email="a@example.test", name="Boss", group="Admin", level="head")
assert m.ob_actor(_admin, {"owner_group": "4W"}) == "Boss (Admin, for 4W)"
assert m.ob_actor(m.User(email="f@example.test", name="Fleet", group="4W", level="staff"), {"owner_group": "4W"}) == "Fleet"
assert m.ob_pickup_moment({"pickup_at": "2026-09-21T10:00"}) == datetime(2026, 9, 21, 10)   # released before
assert m.ob_deadline(NOW, m.ob_pickup_moment(p)) == datetime(2026, 9, 21, 10)
assert m.ob_deadline(datetime(2026, 9, 19), m.ob_pickup_moment(p)) == datetime(2026, 9, 20, 20)
assert m.ob_deadline(datetime(2026, 9, 19), "2026-09-20T10:00") == datetime(2026, 9, 20, 10)
assert "revenue" not in m.ob_payload({**p, "revenue": "999", "_crm": {"price": 999}})
assert "_crm" not in m.ob_payload({**p, "_crm": {"price": 999}})
specs = m.ob_check_specs(p)
owners = {s["check_key"]: s["owner_group"] for s in specs}
assert owners == {"packing:PCK": "CL", "packing:PCK Kayu": "CL", "pickup:fleet": "4W", "pickup:documents": "4W", "pickup:tkbm": "Sort", "delivery:fleet": "2W", "delivery:documents": "2W"}
assert len(m.ob_check_specs({**p, "packing": ["No"]})) == 5
before = {s["check_key"]: s["fingerprint"] for s in specs}
after = {s["check_key"]: s["fingerprint"] for s in m.ob_check_specs({**p, "pod_treatment": "New POD process"})}
assert before["pickup:documents"] != after["pickup:documents"]
assert before["packing:PCK"] == after["packing:PCK"]
cascade = (Path(__file__).parents[1] / "backend/resources/db/migration/V33__operational_onboarding_cascade.sql").read_text(encoding="utf-8")
assert cascade.count("ON DELETE CASCADE") == 5
assert "DELETE FROM" not in cascade
for group in m.OPERATIONAL_GROUPS:
    u = m.User(email="u@example.test", name="Reader", group=group, level="staff")
    assert m.can(u, "operationalOnly") and not m.can(u, "seePrice")
    assert not m.can(u, "editOnboarding", t)
assert m.can(sales, "editOnboarding", t) and not m.can(sales, "approveOnboardingException")
assert m.can(manager, "approveOnboardingException")
assert not m.can(m.User(email="other@example.test", name="Other", group="Commercial", level="staff"), "editOnboarding", t)

class Cur:
    def __init__(self, intake=None, checks=None):
        self.intake = intake
        self.checks = checks or []
        self.sql = []
        self.last = ""
    async def __aenter__(self): return self
    async def __aexit__(self, *args): pass
    async def execute(self, sql, args=()): self.last = sql; self.sql.append((sql, args))
    async def fetchone(self): return self.intake
    async def fetchall(self): return self.checks

class Conn:
    def __init__(self, cur): self.cur = cur; self.committed = False; self.rolled_back = False
    async def __aenter__(self): return self
    async def __aexit__(self, *args): pass
    def cursor(self, *args): return self.cur
    async def begin(self): pass
    async def commit(self): self.committed = True
    async def rollback(self): self.rolled_back = True

class Pool:
    def __init__(self, cur): self.conn = Conn(cur)
    def acquire(self): return self.conn

def checkrow(**kw):
    return dict(id=1, ticket_id=1, fingerprint="fp", label="Pickup fleet", owner_group="4W", status="pending",
                released_payload=json.dumps(p), actual_golive=None, submitted_at=NOW,
                exception_reason=None, feasible_at=None, approved_at=None, confirmed_at=None, **kw)

async def main():
    cur = Cur()
    c = checkrow()
    await denied(lambda: m.ob_apply_decision(cur, c, 1, m.OperationalDecision(fingerprint="fp"), sales), 403)
    await denied(lambda: m.ob_apply_decision(cur, c, 1, m.OperationalDecision(fingerprint="old"), team), 409)
    await m.ob_apply_decision(cur, c, 1, m.OperationalDecision(fingerprint="fp"), team)
    await denied(lambda: m.ob_apply_decision(cur, c, 1, m.OperationalDecision(fingerprint="fp", action="approve"), manager), 409)
    c.update(exception_reason="Fleet unavailable", workaround="Verified replacement fleet", feasible_at=NOW)
    await m.ob_apply_decision(cur, c, 1, m.OperationalDecision(fingerprint="fp", action="approve"), manager)
    assert m.ob_readiness([dict(status="not_ready", approved_at=NOW)]) == "Approved with exception"
    assert m.ob_readiness([dict(status="pending", approved_at=None)]) == "Pending Readiness"
    late = checkrow(); late["submitted_at"] = NOW - timedelta(days=2)
    await denied(lambda: m.ob_apply_decision(cur, late, 1, m.OperationalDecision(fingerprint="fp"), team), 409)
    actual = dict(submitted_at=NOW - timedelta(days=1), actual_golive=None, qc_accepted_at=None, released_payload=json.dumps({**p, "pickup_at":"2026-09-20"}))
    await denied(lambda: m.ob_apply_golive(cur, t, actual, [dict(status="pending")], m.OperationalDate(on="2026-09-20"), sales), 409)
    await m.ob_apply_golive(cur, t, actual, [dict(status="ready", confirmed_at=NOW)], m.OperationalDate(on="2026-09-20"), sales)
    await denied(lambda: m.ob_apply_golive(cur, t, actual, [dict(status="ready", confirmed_at=NOW)], m.OperationalDate(on="2026-09-19"), sales), 400)
    actual["actual_golive"] = NOW.date() - timedelta(days=7)
    await denied(lambda: m.ob_apply_qc(cur, t, actual, qc), 409)
    actual["actual_golive"] -= timedelta(days=1)
    await m.ob_apply_qc(cur, t, actual, qc)
    # Handed over by the Go Live button: QC may accept at once, no seven-day wait
    # (Michael, 2026-09-23).
    handed = dict(submitted_at=NOW, actual_golive=NOW.date(), qc_accepted_at=None, handover_at=NOW)
    await m.ob_apply_qc(cur, t, handed, qc)
    await denied(lambda: m.ob_apply_qc(cur, t, dict(submitted_at=NOW, actual_golive=NOW.date(),
                                                    qc_accepted_at=None), qc), 409)

    async def ticket(ref): return dict(t)
    async def query(sql, args=(), one=False):
        if "onboarding_documents" in sql: return docs
        if "ticket_input" in sql: return {"payload": json.dumps({"revenue": "SECRET_REVENUE", "_crm": {"price": "SECRET_PRICE"},
                                                                  # The charter answers onboarding follows.
                                                                  "commodity": p["product_type"], "shipMode": p["delivery_mode"],
                                                                  "mps": p["mps"], "rdo": p["rdo"], "freq": p["pickup_frequency"]})}
        if "ticket_files" in sql: return []
        if "onboarding_intake" in sql and one: return dict(payload=json.dumps({**p, "revenue": "SECRET_REVENUE"}), released_payload=json.dumps(p), revision=1, submitted_at=NOW, actual_golive=None, qc_accepted_at=None)
        return None if one else []
    async def notice(*args, **kwargs): pass
    m.get_ticket = ticket; m.q = query; m.notify = notice
    for submit in [False, True]:
        savecur = Cur(); m._pool = Pool(savecur)
        await m.operational_save("SOF-1", m.OperationalSave(payload=p, revision=0, submit=submit, documents_reviewed=True), sales)
        writes = [sql for sql, _ in savecur.sql if sql.startswith(("INSERT", "UPDATE", "DELETE"))]
        assert any("operational_master" in sql for sql in writes) == submit
        assert all(not sql.startswith(("UPDATE tickets", "INSERT INTO ticket_input", "UPDATE ticket_input")) for sql in writes)
        assert m._pool.conn.committed
    stale = Cur(intake=dict(revision=2, actual_golive=None)); m._pool = Pool(stale)
    await denied(lambda: m.operational_save("SOF-1", m.OperationalSave(payload=p, revision=0), sales), 409)
    assert m._pool.conn.rolled_back and not m._pool.conn.committed
    safe = await m.operational_detail("SOF-1", team)
    assert "SECRET" not in json.dumps(safe, default=str) and "potential_rev" not in safe

    import httpx
    original = m.current_user
    async def identity(*args): return team
    async def dependency_identity(): return team
    m.current_user = identity
    m.app.dependency_overrides[original] = dependency_identity
    async def rows(): return []
    m.ob_rows = rows
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=m.app), base_url="http://test") as client:
        for path in ["/api/tickets", "/api/stats", "/api/onboarding", "/api/files/1", "/api/users"]:
            assert (await client.get(path)).status_code == 403, path
        assert (await client.get("/api/onboarding-v2")).status_code == 200

asyncio.run(main())
from openpyxl import Workbook
wb = Workbook(); ws = wb.active; ws.append(m.OB_MASTER_HEADERS)
ws.append(["G-1", "PT Test - FTL", "FTL", "4W", "2W"])
buf = io.BytesIO(); wb.save(buf)
assert len(m.ob_excel_rows(buf.getvalue())) == 1
ws["A2"] = 123; buf = io.BytesIO(); wb.save(buf)
rejected(lambda: m.ob_excel_rows(buf.getvalue()), 400)
ws["A2"] = "=1+1"; buf = io.BytesIO(); wb.save(buf)
rejected(lambda: m.ob_excel_rows(buf.getvalue()), 400)
print("ALL OPERATIONAL ONBOARDING CHECKS PASSED")
